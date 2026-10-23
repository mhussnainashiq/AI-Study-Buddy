from pathlib import Path
import io
import fitz
from docx import Document as DocxDocument
from fastapi import HTTPException, UploadFile
from app.core.config import settings

ALLOWED = {".pdf", ".docx", ".txt"}

async def read_upload(file: UploadFile) -> tuple[bytes, str]:
    name = Path(file.filename or "").name
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED:
        raise HTTPException(415, "Unsupported file type. Upload PDF, DOCX, or TXT.")
    data = await file.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if not data:
        raise HTTPException(400, "The uploaded file is empty.")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, f"File exceeds the {settings.max_upload_mb} MB limit.")
    if suffix == ".pdf" and not data.startswith(b"%PDF"):
        raise HTTPException(400, "The file does not appear to be a valid PDF.")
    if suffix == ".docx" and not data.startswith(b"PK"):
        raise HTTPException(400, "The file does not appear to be a valid DOCX document.")
    return data, suffix

def extract_text(data: bytes, suffix: str) -> tuple[str, int]:
    try:
        if suffix == ".txt":
            text = data.decode("utf-8-sig", errors="replace")
            pages = 1
        elif suffix == ".docx":
            doc = DocxDocument(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
            pages = 0
        else:
            pdf = fitz.open(stream=data, filetype="pdf")
            if pdf.is_encrypted:
                raise HTTPException(400, "This PDF is encrypted. Upload an unlocked copy.")
            pages = len(pdf)
            text = "\n\n".join(f"[Page {i+1}]\n{page.get_text()}" for i, page in enumerate(pdf))
            if len(text.strip()) < 40 and settings.ocr_enabled:
                try:
                    import pytesseract
                    from PIL import Image
                    chunks=[]
                    for i, page in enumerate(pdf):
                        pix=page.get_pixmap(matrix=fitz.Matrix(2,2))
                        img=Image.open(io.BytesIO(pix.tobytes("png")))
                        chunks.append(f"[Page {i+1}]\n" + pytesseract.image_to_string(img))
                    text="\n\n".join(chunks)
                except Exception as exc:
                    raise HTTPException(422, "OCR could not process this scanned PDF. Check Tesseract installation.") from exc
            pdf.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, "Could not read this document. It may be corrupted or unsupported.") from exc
    if not text.strip():
        raise HTTPException(422, "No readable text found. For scanned PDFs, enable OCR and install Tesseract.")
    return text[:2_000_000], pages

def chunk_text(text: str, chunk_size: int = 3500, overlap: int = 300) -> list[str]:
    text = text.strip()
    if len(text) <= chunk_size: return [text]
    chunks=[]; start=0
    while start < len(text):
        end=min(start+chunk_size, len(text))
        if end < len(text):
            split=text.rfind("\n", start+chunk_size//2, end)
            if split > start: end=split
        chunks.append(text[start:end])
        if end >= len(text): break
        start=max(end-overlap, start+1)
    return chunks
